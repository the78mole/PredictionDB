"""EIA Short-Term Energy Outlook: Ausgaben Januar 2015 und Oktober 2026 (gekürzt auf die Blätter
``Dates`` und ``2tab``; Ist-Werte von Thomson Reuters durch 999 ersetzt)."""

import io
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from energy_ingest.config import ALL_SOURCES, EnergySettings
from energy_ingest.http import Http
from energy_ingest.model import BERLIN, MONTH, Kind, estimate_published, series
from energy_ingest.sources import FormatError, SourceError
from energy_ingest.sources.steo import EiaSteo, issue_name, parse_xlsx

FIXTURES = Path(__file__).parent / "fixtures"
Handler = Callable[[httpx.Request], httpx.Response]
MakeHttp = Callable[[Handler], Http]
NOV_3 = datetime(2026, 11, 3, 12, tzinfo=UTC)


def _bytes(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def _month(year: int, month: int) -> datetime:
    return datetime(year, month, 1, tzinfo=BERLIN).astimezone(UTC)


def _edit(raw: bytes, sheet: str, row: int, col: int, value: object) -> bytes:
    wb = load_workbook(io.BytesIO(raw))
    wb[sheet].cell(row=row, column=col).value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_issue_name() -> None:
    assert issue_name(2026, 10) == "oct26"
    assert issue_name(2015, 1) == "jan15"


def test_parse_october_2026() -> None:
    out = {d.code: d for d in parse_xlsx(_bytes("steo_oct26_base.xlsx"), 2026, 10)}
    # Tabelle bis Dezember 2027: Vorlauf 0 (Oktober 2026) bis 14
    assert sorted(out) == sorted(
        [f"steo_brent_h{k:02d}" for k in range(15)]
        + [f"steo_henry_hub_h{k:02d}" for k in range(15)]
    )
    ts = _month(2026, 10)
    assert out["steo_brent_h00"].points == [(ts, 111.0)]
    assert out["steo_brent_h14"].points == [(ts, 72.0)]
    assert out["steo_henry_hub_h00"].points[0][1] == pytest.approx(3.15, abs=0.01)
    assert all(d.resolution == MONTH for d in out.values())
    assert out["steo_brent_h03"].info.kind is Kind.OUTLOOK_MONTHLY
    assert out["steo_henry_hub_h03"].info.unit == "USD/MMBtu"
    # Ist-Werte (bis September 2026) werden nicht übernommen
    assert all(v != 999.0 for d in out.values() for _, v in d.points)


def test_parse_january_2015_without_last_historical_month() -> None:
    out = {d.code: d for d in parse_xlsx(_bytes("steo_jan15_base.xlsx"), 2015, 1)}
    # Januar-Ausgabe reicht bis Dezember des Folgejahrs: Vorlauf 0 bis 23
    assert len(out) == 48
    assert "steo_brent_h23" in out and "steo_brent_h24" not in out
    assert out["steo_brent_h00"].points == [(_month(2015, 1), 49.0)]
    assert all(v != 999.0 for d in out.values() for _, v in d.points)


def test_wrong_issue_month_is_format_error() -> None:
    with pytest.raises(FormatError, match="Ausgabemonat"):
        parse_xlsx(_bytes("steo_oct26_base.xlsx"), 2026, 9)


@pytest.mark.parametrize(
    ("row", "col", "value", "match"),
    [
        (7, 1, "BRENT_NEU", "keine Prognosen"),  # Code umbenannt
        (5, 2, "Crude Oil (euro per barrel)", "Einheit"),
        (4, 3, "Monat", "Kopfzeile"),
        (4, 8, "6", "lückenlos"),
    ],
)
def test_format_changes(row: int, col: int, value: str, match: str) -> None:
    raw = _edit(_bytes("steo_oct26_base.xlsx"), "2tab", row, col, value)
    with pytest.raises(FormatError, match=match):
        parse_xlsx(raw, 2026, 10)


def test_not_an_xlsx() -> None:
    with pytest.raises(FormatError):
        parse_xlsx(b"<html>Fehler</html>", 2026, 10)


def _source(make_http: MakeHttp, calls: list[str]) -> EiaSteo:
    files = {"oct26": "steo_oct26_base.xlsx", "jan15": "steo_jan15_base.xlsx"}

    def handler(req: httpx.Request) -> httpx.Response:
        calls.append(req.url.path)
        name = req.url.path.rsplit("/", 1)[-1].removesuffix("_base.xlsx")
        if name in files:
            return httpx.Response(200, content=_bytes(files[name]))
        return httpx.Response(404)

    return EiaSteo(make_http(handler))


def test_fetch_skips_unpublished_current_issue(make_http: MakeHttp) -> None:
    calls: list[str] = []
    src = _source(make_http, calls)
    now = NOV_3  # Novemberausgabe noch nicht erschienen
    out = src.fetch("outlook", _month(2026, 10), _month(2026, 12), now=now)
    assert {d.points[0][0] for d in out} == {_month(2026, 10)}
    assert calls == [
        "/outlooks/steo/archives/oct26_base.xlsx",
        "/outlooks/steo/archives/nov26_base.xlsx",
    ]
    # zweiter Abgleich: Oktober aus dem Speicher, November erneut versucht
    src.fetch("outlook", _month(2026, 10), _month(2026, 12), now=now)
    assert calls[2:] == ["/outlooks/steo/archives/nov26_base.xlsx"]


def test_fetch_missing_old_issue_is_error(make_http: MakeHttp) -> None:
    src = _source(make_http, [])
    with pytest.raises(SourceError, match="sep26"):
        src.fetch(
            "outlook", _month(2026, 9), _month(2026, 10), now=datetime(2026, 11, 3, tzinfo=UTC)
        )


def test_fetch_window_uses_issue_months(make_http: MakeHttp) -> None:
    calls: list[str] = []
    src = _source(make_http, calls)
    # Fenster mitten im Monat: nur Ausgaben, deren Monatsbeginn im Fenster liegt
    start = datetime(2014, 12, 20, tzinfo=UTC)
    out = src.fetch("outlook", start, _month(2015, 1) + timedelta(days=3))
    assert calls == ["/outlooks/steo/archives/jan15_base.xlsx"]
    assert len(out) == 48


def test_published_estimate() -> None:
    ts = _month(2026, 10)
    assert estimate_published(Kind.OUTLOOK_MONTHLY, ts, MONTH) == ts + timedelta(days=14)
    assert series("steo_brent_h23").unit == "USD/bbl"


def test_steo_is_free_and_active_by_default() -> None:
    assert "eia_steo" in ALL_SOURCES
    assert EnergySettings().allowed("eia_steo")
    assert EiaSteo.free
