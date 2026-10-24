"""Weltbank, Commodity Price Data ("Pink Sheet"): Monatsmittel Öl, Kohle, Erdgas, LNG.

Datei ``CMO-Historical-Data-Monthly.xlsx``, Blatt ``Monthly Prices``, Monatswerte ab 1960
in nominalen US-Dollar, Lizenz CC BY 4.0. Erdgas Europa entspricht TTF, Kohle Australien
Newcastle, Kohle Südafrika Richards Bay.

Die Weltbank legt die Datei jedes Jahr unter einer neuen Adresse ab (die Dokument-ID in der
URL enthält das Jahr). Die Adresse ist deshalb einstellbar (``PDB_ENERGY_WORLDBANK_URL``);
eine veraltete Datei fällt im Grafana-Alarm "Reihe veraltet" auf.

Monatswerte stehen am Monatsbeginn (00:00 Ortszeit, als UTC).
"""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from typing import Any, ClassVar

from openpyxl import load_workbook

from energy_ingest.http import Http
from energy_ingest.model import BERLIN, MONTH, SeriesData, series
from energy_ingest.sources import Cached, FormatError, SourceError, clip

URL = (
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/"
    "related/CMO-Historical-Data-Monthly.xlsx"
)
SHEET = "Monthly Prices"
TTL = timedelta(hours=12)

# Spaltenkopf -> (Reihen-Code, Einheitenzeile)
COLUMNS: dict[str, tuple[str, str]] = {
    "Crude oil, Brent": ("oil_brent_m", "($/bbl)"),
    "Crude oil, WTI": ("oil_wti_m", "($/bbl)"),
    "Crude oil, Dubai": ("oil_dubai_m", "($/bbl)"),
    "Coal, Australian": ("coal_au_m", "($/mt)"),
    "Coal, South African **": ("coal_za_m", "($/mt)"),
    "Natural gas, US": ("gas_henry_hub_m", "($/mmbtu)"),
    "Natural gas, Europe": ("gas_eu_m", "($/mmbtu)"),
    "Liquefied natural gas, Japan": ("lng_japan_m", "($/mmbtu)"),
}


def _month(text: str) -> datetime | None:
    """``1960M01`` -> Monatsbeginn; andere Zellen (Fußnoten) -> None."""
    if len(text) != 7 or text[4] != "M" or not (text[:4] + text[5:]).isdigit():
        return None
    return datetime(int(text[:4]), int(text[5:]), 1, tzinfo=BERLIN).astimezone(UTC)


def _value(cell: Any) -> float | None:
    return float(cell) if isinstance(cell, int | float) else None


def parse_xlsx(raw: bytes) -> list[SeriesData]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise FormatError(f"Weltbank: keine lesbare xlsx-Datei ({exc})") from exc
    if SHEET not in wb.sheetnames:
        raise FormatError(f"Weltbank: Blatt {SHEET!r} fehlt, vorhanden {wb.sheetnames}")
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()
    head = next((i for i, r in enumerate(rows) if "Crude oil, Brent" in r), None)
    if head is None or head + 1 >= len(rows):
        raise FormatError("Weltbank: Kopfzeile mit 'Crude oil, Brent' fehlt")
    header = [str(c).strip() if c is not None else "" for c in rows[head]]
    units = rows[head + 1]
    cols: dict[str, int] = {}
    for name, (code, unit) in COLUMNS.items():
        if name not in header:
            raise FormatError(f"Weltbank: Spalte {name!r} fehlt")
        i = header.index(name)
        if str(units[i]).strip() != unit:
            raise FormatError(f"Weltbank: Einheit {units[i]!r} für {name!r}, erwartet {unit!r}")
        cols[code] = i
    points: dict[str, list[tuple[datetime, float]]] = {code: [] for code in cols}
    for row in rows[head + 2 :]:
        ts = _month(str(row[0] or "").strip())
        if ts is None:
            continue
        for code, i in cols.items():
            v = _value(row[i])
            if v is not None:
                points[code].append((ts, v))
    if not any(points.values()):
        raise FormatError("Weltbank: keine Monatswerte gefunden")
    return [SeriesData(series(code), MONTH, pts) for code, pts in points.items() if pts]


class WorldBank:
    code = "worldbank"
    license = "CC BY 4.0, World Bank Commodity Price Data (The Pink Sheet)"
    priority = 10
    free = True
    # Der Monatswert erscheint Anfang des Folgemonats
    sync_lookback: ClassVar[dict[str, timedelta]] = {"prices": timedelta(days=95)}

    def __init__(self, http: Http, url: str = URL) -> None:
        self.http = http
        self.url = url
        self._data = Cached(self._load, TTL)

    def _load(self) -> list[SeriesData]:
        r = self.http.get(self.url)
        if r.status_code >= 400:
            raise SourceError(f"Weltbank: {r.status_code} für {self.url}")
        return parse_xlsx(r.content)

    def datasets(self) -> list[str]:
        return ["prices"]

    def max_chunk(self, dataset: str) -> timedelta:
        return timedelta(days=3650)

    def fetch(self, dataset: str, start: datetime, end: datetime) -> list[SeriesData]:
        if dataset != "prices":
            raise KeyError(dataset)
        out = []
        for data in self._data.get():
            pts = clip(data.points, start, end)
            if pts:
                out.append(SeriesData(data.info, data.resolution, pts))
        return out
