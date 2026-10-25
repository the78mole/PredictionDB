"""EIA Short-Term Energy Outlook (STEO): Monatsprognosen Brent und Henry Hub je Ausgabe.

Die EIA veröffentlicht den STEO monatlich (meist zwischen dem 6. und 12.) und legt jede Ausgabe
unter einer festen Adresse ins Archiv: ``https://www.eia.gov/outlooks/steo/archives/
<mmmJJ>_base.xlsx`` (z. B. ``oct26_base.xlsx``), auch die jeweils aktuelle. Gelesen wird das
Blatt ``2tab`` (Energy Prices), Zeilen ``BREPUUS`` (Brent Spot Average, USD/bbl) und
``NGHHUUS`` (Henry Hub Spot, USD/MMBtu), dazu das Blatt ``Dates`` (Ausgabemonat und letzter
Ist-Monat).

Nutzungsrecht: Die Prognosen sind Werk der EIA und gemeinfrei
(https://www.eia.gov/about/copyrights_reuse.php). Die Ist-Werte von Brent und Henry Hub in
derselben Tabelle stammen von Thomson Reuters (vgl. Quelle ``eia_reuters``); sie werden nicht
übernommen. Gespeichert werden nur Monate nach dem letzten Ist-Monat, also nur Prognosen.

Je Ausgabe und Vorlauf k entsteht ein Wert der Reihe ``steo_<brent|henry_hub>_h<kk>`` mit dem
Zeitstempel des Ausgabemonats (Monatsbeginn, 00:00 Ortszeit, als UTC) und der Prognose für den
Monat Ausgabemonat + k. ``h00`` ist die Schätzung für den laufenden Monat.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from functools import partial
from typing import Any, ClassVar

from openpyxl import load_workbook

from energy_ingest.http import Http
from energy_ingest.model import (
    BERLIN,
    MONTH,
    OUTLOOK_HORIZONS,
    SeriesData,
    next_period,
    outlook_code,
    series,
)
from energy_ingest.sources import Cached, FormatError, SourceError

URL = "https://www.eia.gov/outlooks/steo/archives/{name}_base.xlsx"
SHEET = "2tab"
TTL = timedelta(hours=12)
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
MONTH_NAMES = [
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
]
# STEO-Code -> Kürzel in ``OUTLOOK_SERIES``
CODES = {"BREPUUS": "brent", "NGHHUUS": "henry_hub"}
# erwarteter Text der Einheit in der Zeile bzw. im Abschnitt darüber
UNITS = {"BREPUUS": "dollars per barrel", "NGHHUUS": "dollars per million btu"}
FIRST_ISSUE = datetime(2015, 1, 1, tzinfo=BERLIN)


def issue_name(year: int, month: int) -> str:
    """Dateiname einer Ausgabe ohne Endung, z. B. ``oct26``."""
    return f"{MONTHS[month - 1]}{year % 100:02d}"


def _month_index(year: int, month: int) -> int:
    return year * 12 + month - 1


def _issue_of(text: Any) -> int | None:
    """``October 2026`` -> Monatsindex."""
    parts = str(text or "").strip().lower().split()
    if len(parts) != 2 or parts[0] not in MONTH_NAMES or not parts[1].isdigit():
        return None
    return _month_index(int(parts[1]), MONTH_NAMES.index(parts[0]) + 1)


def _value(cell: Any) -> float | None:
    return float(cell) if isinstance(cell, int | float) and cell > 0 else None


def parse_xlsx(raw: bytes, year: int, month: int) -> list[SeriesData]:
    """Prognosen einer Ausgabe als Reihen je Vorlauf (je ein Wert am Ausgabemonat)."""
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise FormatError(
            f"STEO {issue_name(year, month)}: keine lesbare xlsx-Datei ({exc})"
        ) from exc
    name = issue_name(year, month)
    for sheet in ("Dates", SHEET):
        if sheet not in wb.sheetnames:
            raise FormatError(f"STEO {name}: Blatt {sheet!r} fehlt")
    dates = {
        str(r[0]).strip().rstrip("- ").lower(): r[3]
        for r in wb["Dates"].iter_rows(max_row=20, values_only=True)
        if r and r[0] is not None and len(r) > 3
    }
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()

    issue = _month_index(year, month)
    found = _issue_of(dates.get("forecast month"))
    if found != issue:
        raise FormatError(f"STEO {name}: Ausgabemonat {dates.get('forecast month')!r}")
    # ältere Ausgaben nennen den letzten Ist-Monat nicht; er ist dann der Vormonat
    last = dates.get("last historical month")
    last_hist = (
        _month_index(int(last) // 100, int(last) % 100) if isinstance(last, int) else issue - 1
    )

    # Kopf: Zeile mit Jahreszahlen (nur bei Januar), darunter Zeile mit Monatskürzeln
    head = next(
        (i for i, r in enumerate(rows[:10]) if [str(c)[:3].lower() for c in r[2:5]] == MONTHS[:3]),
        None,
    )
    if head is None or head == 0:
        raise FormatError(f"STEO {name}: Kopfzeile mit Monaten fehlt")
    cols: dict[int, int] = {}  # Spalte -> Monatsindex
    year_of = None
    for j, (y, m) in enumerate(zip(rows[head - 1], rows[head], strict=False)):
        if j < 2:
            continue
        if isinstance(y, int):
            year_of = y
        mon = str(m or "")[:3].lower()
        if year_of is None or mon not in MONTHS:
            continue
        cols[j] = _month_index(year_of, MONTHS.index(mon) + 1)
    order = list(cols.values())
    if not order or order != list(range(order[0], order[0] + len(order))):
        raise FormatError(f"STEO {name}: Monatsspalten nicht lückenlos")

    ts = datetime(year, month, 1, tzinfo=BERLIN).astimezone(UTC)
    out = []
    seen: set[str] = set()
    section = ""
    for r in rows[head + 1 :]:
        label = str(r[1] or "").strip().lower() if len(r) > 1 else ""
        if r[0] is None and label:
            section = label
        code = str(r[0] or "").strip()
        if code not in CODES:
            continue
        if UNITS[code] not in f"{section} {label}":
            raise FormatError(f"STEO {name}: Einheit für {code} weicht ab ({label!r})")
        for j, idx in cols.items():
            k = idx - issue
            if idx <= last_hist or k not in OUTLOOK_HORIZONS:
                continue
            v = _value(r[j]) if j < len(r) else None
            if v is not None:
                out.append(SeriesData(series(outlook_code(CODES[code], k)), MONTH, [(ts, v)]))
                seen.add(CODES[code])
    missing = set(CODES.values()) - seen
    if missing:
        raise FormatError(f"STEO {name}: keine Prognosen für {sorted(missing)}")
    return out


class EiaSteo:
    code = "eia_steo"
    license = (
        "Public Domain, U.S. Energy Information Administration (EIA), Short-Term Energy Outlook"
    )
    priority = 10
    free = True
    # Abgleich: die beiden letzten Ausgaben, falls eine später erschienen ist
    sync_lookback: ClassVar[dict[str, timedelta]] = {"outlook": timedelta(days=62)}

    def __init__(self, http: Http, url: str = URL) -> None:
        self.http = http
        self.url = url
        self._issues: dict[str, Cached[list[SeriesData] | None]] = {}

    def datasets(self) -> list[str]:
        return ["outlook"]

    def max_chunk(self, dataset: str) -> timedelta:
        return timedelta(days=366)

    def _load(self, year: int, month: int, now: datetime) -> list[SeriesData] | None:
        url = self.url.format(name=issue_name(year, month))
        r = self.http.get(url)
        if r.status_code == 404:
            # die Ausgabe des laufenden Monats erscheint erst in dessen erster Hälfte
            local = now.astimezone(BERLIN)
            if (year, month) >= (local.year, local.month):
                return None
            raise SourceError(f"STEO: Ausgabe {issue_name(year, month)} fehlt ({url})")
        if r.status_code >= 400:
            raise SourceError(f"STEO: {r.status_code} für {url}")
        return parse_xlsx(r.content, year, month)

    def fetch(
        self, dataset: str, start: datetime, end: datetime, now: datetime | None = None
    ) -> list[SeriesData]:
        if dataset != "outlook":
            raise KeyError(dataset)
        now = now or datetime.now(UTC)
        out: list[SeriesData] = []
        ts = max(start, FIRST_ISSUE.astimezone(UTC))
        local = ts.astimezone(BERLIN)
        ts = datetime(local.year, local.month, 1, tzinfo=BERLIN).astimezone(UTC)
        if ts < start:
            ts = next_period(ts, 1)
        stop = min(end, now)
        while ts < stop:
            local = ts.astimezone(BERLIN)
            name = issue_name(local.year, local.month)
            cached = self._issues.get(name)
            if cached is None:
                load = partial(self._load, local.year, local.month, now)
                cached = self._issues[name] = Cached(load, TTL)
            data = cached.get()
            if data is None:
                del self._issues[name]  # noch nicht erschienen: beim nächsten Abgleich erneut
            else:
                out.extend(data)
            ts = next_period(ts, 1)
        return out
