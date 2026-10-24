"""EU Weekly Oil Bulletin (Europäische Kommission, GD Energie): Endkundenpreise Deutschland.

Historien-Datei ``Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx`` mit den Blättern
``Prices with taxes`` und ``Prices wo taxes``; je Land sechs Spalten (Euro-super 95, Diesel,
Heizöl, zwei Sorten schweres Heizöl, LPG), wöchentlich ab 03.01.2005, Stichtag Montag.
Lizenz CC BY 4.0 (Beschluss 2011/833/EU über die Weiterverwendung von Kommissionsdokumenten).

Die Download-Adresse enthält eine Dokument-ID und ist deshalb einstellbar
(``PDB_ENERGY_OIL_BULLETIN_URL``). Weicht der Aufbau ab (Blatt, Spalte, Einheit), schreibt
die Quelle nichts und meldet ``status = 'format'``.

Werte in EUR je Liter (Datei: EUR je 1000 l), am Stichtag 00:00 Ortszeit (als UTC).
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from typing import ClassVar

from openpyxl import load_workbook

from energy_ingest.http import Http
from energy_ingest.model import WEEK, SeriesData, series
from energy_ingest.sources import Cached, FormatError, SourceError, clip, day_start

URL = (
    "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)
TTL = timedelta(hours=12)
COUNTRY = "DE"

# Blatt -> (Spaltenpräfix, Suffix des Reihen-Codes)
SHEETS = {
    "Prices with taxes": ("price_with_tax", ""),
    "Prices wo taxes": ("price_wo_tax", "_net"),
}
# Produkt in der Datei -> Produkt im Reihen-Code
PRODUCTS = {"euro95": "euro95", "diesel": "diesel", "heating_oil": "heating_oil", "LPG": "lpg"}
UNIT = "1000 l"


def parse_sheet(rows: list[tuple[object, ...]], prefix: str, suffix: str) -> list[SeriesData]:
    if len(rows) < 4:
        raise FormatError(f"Oil Bulletin: Blatt mit Präfix {prefix!r} zu kurz")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    units = rows[2]
    cols: dict[str, int] = {}
    for product, code in PRODUCTS.items():
        name = f"{COUNTRY}_{prefix}_{product}"
        if name not in header:
            raise FormatError(f"Oil Bulletin: Spalte {name!r} fehlt")
        i = header.index(name)
        if str(units[i]).strip() != UNIT:
            raise FormatError(f"Oil Bulletin: Einheit {units[i]!r} in {name!r}, erwartet {UNIT!r}")
        cols[f"retail_de_{code}{suffix}"] = i
    points: dict[str, list[tuple[datetime, float]]] = {code: [] for code in cols}
    for row in rows[3:]:
        if not isinstance(row[0], datetime):
            continue  # Leer- und Fußnotenzeilen
        ts = day_start(row[0].date())
        for code, i in cols.items():
            v = row[i]
            if isinstance(v, int | float):
                points[code].append((ts, float(v) / 1000.0))
    if not any(points.values()):
        raise FormatError(f"Oil Bulletin: keine Werte für {COUNTRY} ({prefix})")
    return [SeriesData(series(c), WEEK, sorted(pts)) for c, pts in points.items() if pts]


def parse_xlsx(raw: bytes) -> list[SeriesData]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise FormatError(f"Oil Bulletin: keine lesbare xlsx-Datei ({exc})") from exc
    out = []
    for sheet, (prefix, suffix) in SHEETS.items():
        if sheet not in wb.sheetnames:
            raise FormatError(f"Oil Bulletin: Blatt {sheet!r} fehlt, vorhanden {wb.sheetnames}")
        out.extend(parse_sheet(list(wb[sheet].iter_rows(values_only=True)), prefix, suffix))
    wb.close()
    return out


class OilBulletin:
    code = "oil_bulletin"
    license = "CC BY 4.0, Europäische Kommission, Weekly Oil Bulletin"
    priority = 10
    free = True
    sync_lookback: ClassVar[dict[str, timedelta]] = {"retail_de": timedelta(days=21)}

    def __init__(self, http: Http, url: str = URL) -> None:
        self.http = http
        self.url = url
        self._data = Cached(self._load, TTL)

    def _load(self) -> list[SeriesData]:
        r = self.http.get(self.url)
        if r.status_code >= 400:
            raise SourceError(f"Oil Bulletin: {r.status_code} für {self.url}")
        return parse_xlsx(r.content)

    def datasets(self) -> list[str]:
        return ["retail_de"]

    def max_chunk(self, dataset: str) -> timedelta:
        return timedelta(days=3650)

    def fetch(self, dataset: str, start: datetime, end: datetime) -> list[SeriesData]:
        if dataset != "retail_de":
            raise KeyError(dataset)
        out = []
        for data in self._data.get():
            pts = clip(data.points, start, end)
            if pts:
                out.append(SeriesData(data.info, data.resolution, pts))
        return out
